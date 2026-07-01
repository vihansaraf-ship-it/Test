"""Excel report generation with formatted sheets.

Handles both PSE board reports and simple income statements.
"""

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FILL_GREEN = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
FILL_RED = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")


def _auto_width(ws):
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 30)


def _style_headers(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def write_excel(df: pd.DataFrame, output_path: str) -> None:
    is_pse = "ebitda_ytd" in df.columns
    if is_pse:
        _write_pse_excel(df, output_path)
    else:
        _write_simple_excel(df, output_path)


def _write_pse_excel(df: pd.DataFrame, output_path: str):
    df = df.sort_values("period")

    # Summary sheet: key KPIs per period
    summary_cols = [
        "period", "revenue", "revenue_budget",
        "total_revenue", "total_revenue_budget",
        "opex", "opex_budget",
        "total_ga", "total_ga_budget",
        "ebitda", "ebitda_budget",
        "ebitda_margin_pct",
    ]
    summary_cols = [c for c in summary_cols if c in df.columns]

    # YTD sheet
    ytd_cols = [
        "period",
        "revenue_ytd", "revenue_ytd_budget",
        "total_revenue_ytd", "total_revenue_ytd_budget",
        "opex_ytd", "opex_ytd_budget",
        "total_ga_ytd", "total_ga_ytd_budget",
        "ebitda_ytd", "ebitda_ytd_budget",
        "ebitda_margin_ytd_pct",
    ]
    ytd_cols = [c for c in ytd_cols if c in df.columns]

    # Rename columns for readability
    col_names = {
        "period": "Period",
        "revenue": "Revenue (MTD)", "revenue_budget": "Revenue Budget",
        "total_revenue": "Total Revenue (MTD)", "total_revenue_budget": "Total Rev Budget",
        "opex": "OPEX (MTD)", "opex_budget": "OPEX Budget",
        "total_ga": "G&A (MTD)", "total_ga_budget": "G&A Budget",
        "ebitda": "EBITDA (MTD)", "ebitda_budget": "EBITDA Budget",
        "ebitda_margin_pct": "EBITDA Margin %",
        "revenue_ytd": "Revenue (YTD)", "revenue_ytd_budget": "Revenue Budget (YTD)",
        "total_revenue_ytd": "Total Revenue (YTD)", "total_revenue_ytd_budget": "Total Rev Budget (YTD)",
        "opex_ytd": "OPEX (YTD)", "opex_ytd_budget": "OPEX Budget (YTD)",
        "total_ga_ytd": "G&A (YTD)", "total_ga_ytd_budget": "G&A Budget (YTD)",
        "ebitda_ytd": "EBITDA (YTD)", "ebitda_ytd_budget": "EBITDA Budget (YTD)",
        "ebitda_margin_ytd_pct": "EBITDA Margin % (YTD)",
    }

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # MTD Summary
        summary_df = df[summary_cols].rename(columns=col_names)
        summary_df.to_excel(writer, sheet_name="MTD Summary", index=False)
        ws = writer.sheets["MTD Summary"]
        _style_headers(ws)
        _format_currency_cols(ws, ["Revenue (MTD)", "Revenue Budget", "Total Revenue (MTD)",
                                    "Total Rev Budget", "OPEX (MTD)", "OPEX Budget",
                                    "G&A (MTD)", "G&A Budget", "EBITDA (MTD)", "EBITDA Budget"])
        _format_pct_cols(ws, ["EBITDA Margin %"])
        _apply_ebitda_coloring(ws, "EBITDA (MTD)")
        _auto_width(ws)

        # YTD Summary
        ytd_df = df[ytd_cols].rename(columns=col_names)
        ytd_df.to_excel(writer, sheet_name="YTD Summary", index=False)
        ws = writer.sheets["YTD Summary"]
        _style_headers(ws)
        _format_currency_cols(ws, ["Revenue (YTD)", "Revenue Budget (YTD)", "Total Revenue (YTD)",
                                    "Total Rev Budget (YTD)", "OPEX (YTD)", "OPEX Budget (YTD)",
                                    "G&A (YTD)", "G&A Budget (YTD)", "EBITDA (YTD)", "EBITDA Budget (YTD)"])
        _format_pct_cols(ws, ["EBITDA Margin % (YTD)"])
        _auto_width(ws)

        # Full Detail
        detail_cols = ["company", "period", "source_file"] + [
            c for c in df.columns
            if c not in ("company", "period", "source_file", "unit")
        ]
        detail_cols = [c for c in detail_cols if c in df.columns]
        df[detail_cols].to_excel(writer, sheet_name="Full Detail", index=False)
        ws = writer.sheets["Full Detail"]
        _style_headers(ws)
        _auto_width(ws)


def _write_simple_excel(df: pd.DataFrame, output_path: str):
    summary_df = df.sort_values("period").groupby("company").last().reset_index()
    summary_cols = ["company", "period", "revenue", "ebitda", "dscr"]
    summary_cols = [c for c in summary_cols if c in summary_df.columns]
    summary_df = summary_df[summary_cols]

    detail_cols = [
        "company", "period", "source_file", "revenue", "cogs", "gross_profit",
        "opex", "depreciation_amortization", "ebit", "interest_expense",
        "tax_expense", "net_income", "ebitda", "total_debt_service", "dscr",
    ]
    detail_cols = [c for c in detail_cols if c in df.columns]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.sheets["Summary"]
        _style_headers(ws)
        _format_currency_cols(ws, ["revenue", "ebitda"])
        if "dscr" in summary_cols:
            _apply_dscr_formatting(ws, summary_cols.index("dscr") + 1)
        _auto_width(ws)

        df[detail_cols].to_excel(writer, sheet_name="Detail", index=False)
        ws = writer.sheets["Detail"]
        _style_headers(ws)
        _auto_width(ws)


def _format_currency_cols(ws, col_names: list[str]):
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    for name in col_names:
        col_idx = headers.get(name)
        if not col_idx:
            continue
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell = row[0]
            if cell.value is not None:
                cell.number_format = '#,##0'


def _format_pct_cols(ws, col_names: list[str]):
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    for name in col_names:
        col_idx = headers.get(name)
        if not col_idx:
            continue
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell = row[0]
            if cell.value is not None:
                cell.number_format = '0.0"%"'


def _apply_ebitda_coloring(ws, col_name: str):
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    col_idx = headers.get(col_name)
    if not col_idx:
        return
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        if cell.value is None:
            continue
        try:
            val = float(cell.value)
        except (ValueError, TypeError):
            continue
        cell.fill = FILL_GREEN if val > 0 else FILL_RED


def _apply_dscr_formatting(ws, col_idx: int):
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        if cell.value is None:
            continue
        try:
            val = float(cell.value)
        except (ValueError, TypeError):
            continue
        if val >= 1.5:
            cell.fill = FILL_GREEN
        elif val >= 1.2:
            cell.fill = FILL_YELLOW
        else:
            cell.fill = FILL_RED
